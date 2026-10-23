#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import json
import os
import sqlite3
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

try:
    from init_db import seed_antibiotics
except ImportError:
    from scripts.init_db import seed_antibiotics

try:
    from episode_model import migrate_episode_model
except ImportError:
    from scripts.episode_model import migrate_episode_model


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SHEET = "OrthoDok Knie ges."
NULL_MARKERS = {"", "#", "n.e.", "n.a.", "na", "none", "null"}
DIRECT_IDENTIFIER_HEADERS = {
    "Patient",
    "Geburtsdatum",
    "Nachname",
    "Vorname",
    "Fall",
    "ISH: Mediz. Fall",
    "Medizinischer Fall",
    "Med. Fallnr.",
}


def clean(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        if value.lower() in NULL_MARKERS:
            return None
    return value


def as_text(value: Any) -> str | None:
    value = clean(value)
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def as_iso_date(value: Any) -> str | None:
    value = clean(value)
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return text


def year_from_date(value: str | None) -> int | None:
    if not value:
        return None
    try:
        return int(value[:4])
    except ValueError:
        return None


def file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def stable_hash(*parts: Any, salt: str) -> str:
    raw = "|".join("" if p is None else str(p) for p in parts)
    return hashlib.sha256(f"{salt}|{raw}".encode("utf-8")).hexdigest()[:32]


def dict_from_row(headers: list[str], values: tuple[Any, ...]) -> dict[str, Any]:
    row: dict[str, Any] = {}
    for i, header in enumerate(headers):
        value = values[i] if i < len(values) else None
        if isinstance(value, (datetime, date)):
            value = as_iso_date(value)
        row[header] = value
    return row


def normalized_headers(values: tuple[Any, ...]) -> list[str]:
    seen: dict[str, int] = {}
    headers: list[str] = []
    for index, value in enumerate(values, start=1):
        header = as_text(value) or f"Unnamed_{index}"
        if header in seen:
            seen[header] += 1
            header = f"{header}_{seen[header]}"
        else:
            seen[header] = 1
        headers.append(header)
    return headers


def infer_side(value: Any) -> str | None:
    text = as_text(value)
    if not text:
        return None
    lower = text.lower()
    if lower in {"rechts", "right", "r", "1"}:
        return "rechts"
    if lower in {"links", "left", "l", "2"}:
        return "links"
    return text


def infer_intervention_type(version: str | None, row: dict[str, Any]) -> tuple[int, str]:
    text = " ".join(
        as_text(v) or ""
        for v in [
            version,
            row.get("Intervention - Totalknieendoprothese"),
            row.get("Intervention - Totale Hüftendoprothese"),
            row.get("Intervention - Totale Hueftendoprothese"),
            row.get("Intervention - Beidseitige Prothese"),
        ]
    )
    lower = text.lower()
    if "erst" in lower or "primaer" in lower or "primar" in lower:
        return 1, "Primaerimplantation"
    if "dair" in lower:
        return 2, "DAIR"
    if "debrid" in lower:
        return 3, "Debridement"
    if "spacer" in lower:
        return 4, "Spacer"
    if "explant" in lower:
        return 5, "Explantation"
    if "reimplant" in lower:
        return 6, "Reimplantation"
    if "revision" in lower or "wechsel" in lower:
        return 7, "Revision"
    return 99, version or "Sonstige"


def get_or_create_source(con: sqlite3.Connection, source_path: Path) -> int:
    digest = file_sha256(source_path)
    row = con.execute(
        "SELECT SourceFileID FROM tblSourceFile WHERE FileHashSHA256 = ?",
        (digest,),
    ).fetchone()
    if row:
        return int(row[0])
    cur = con.execute(
        """
        INSERT INTO tblSourceFile (SourceSystem, OriginalPath, FileName, FileHashSHA256)
        VALUES ('Orthodok', ?, ?, ?)
        """,
        (str(source_path), source_path.name, digest),
    )
    return int(cur.lastrowid)


def get_or_create_patient(
    con: sqlite3.Connection,
    row: dict[str, Any],
    *,
    include_identifiers: bool,
    salt: str,
) -> int:
    patient_no = as_text(row.get("Patient"))
    fall_no = as_text(row.get("Fall"))
    birth = as_iso_date(row.get("Geburtsdatum"))
    sex = as_text(row.get("Geschlecht"))
    last_name = as_text(row.get("Nachname"))
    first_name = as_text(row.get("Vorname"))
    if patient_no:
        patient_key = f"patient:{patient_no}"
    elif any([last_name, first_name, birth]):
        patient_key = f"name_birth:{last_name}|{first_name}|{birth}"
    else:
        patient_key = f"fall:{fall_no}"
    patient_hash = stable_hash(patient_key, salt=salt)

    found = con.execute(
        "SELECT PatientID FROM tblPatient WHERE PatientHash = ?",
        (patient_hash,),
    ).fetchone()
    if found:
        return int(found[0])

    cur = con.execute(
        """
        INSERT INTO tblPatient (
            PatientenzahlSAP, PatientHash, Nachname, Vorname, Geburtsdatum,
            Geburtsjahr, Geschlecht, Quelle
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, 'Orthodok')
        """,
        (
            patient_no if include_identifiers else None,
            patient_hash,
            last_name if include_identifiers else None,
            first_name if include_identifiers else None,
            birth if include_identifiers else None,
            year_from_date(birth),
            sex,
        ),
    )
    return int(cur.lastrowid)


def get_or_create_joint_case(
    con: sqlite3.Connection,
    patient_id: int,
    joint: str,
    side: str | None,
    op_date: str | None,
    intervention_type_id: int,
) -> int:
    row = con.execute(
        """
        SELECT GelenkFallID, PrimaerimplantationDatum
        FROM tblGelenkFall
        WHERE PatientID = ? AND Gelenk = ? AND Seite IS ?
        """,
        (patient_id, joint, side),
    ).fetchone()
    if row:
        gelenk_fall_id = int(row[0])
        current_primary = row[1]
        if intervention_type_id == 1 and op_date and (current_primary is None or op_date < current_primary):
            con.execute(
                """
                UPDATE tblGelenkFall
                SET PrimaerimplantationDatum = ?, UpdatedAt = CURRENT_TIMESTAMP
                WHERE GelenkFallID = ?
                """,
                (op_date, gelenk_fall_id),
            )
        return gelenk_fall_id

    cur = con.execute(
        """
        INSERT INTO tblGelenkFall (PatientID, Gelenk, Seite, PrimaerimplantationDatum, Quelle)
        VALUES (?, ?, ?, ?, 'Orthodok')
        """,
        (patient_id, joint, side, op_date if intervention_type_id == 1 else None),
    )
    return int(cur.lastrowid)


def get_or_create_episode(
    con: sqlite3.Connection,
    gelenk_fall_id: int,
    op_date: str | None,
    intervention_type_id: int,
    intervention_text: str,
) -> int:
    found = con.execute(
        """
        SELECT EpisodeID
        FROM tblEpisode
        WHERE GelenkFallID = ?
        ORDER BY EpisodeBeginn IS NULL, EpisodeBeginn, EpisodeID
        LIMIT 1
        """,
        (gelenk_fall_id,),
    ).fetchone()
    if found:
        return int(found[0])

    episode_type = "Primaer-/Registerepisode" if intervention_type_id == 1 else "PJI"
    cur = con.execute(
        """
        INSERT INTO tblEpisode (
            GelenkFallID, EpisodeTyp, EpisodeBeginn, Kommentar, Quelle
        )
        VALUES (?, ?, ?, ?, 'Orthodok')
        """,
        (
            gelenk_fall_id,
            episode_type,
            op_date,
            f"Automatisch aus OrthoDok-Eingriff angelegt: {intervention_text}",
        ),
    )
    return int(cur.lastrowid)


def get_or_create_stay(
    con: sqlite3.Connection,
    gelenk_fall_id: int,
    episode_id: int,
    row: dict[str, Any],
    *,
    include_identifiers: bool,
    salt: str,
) -> int:
    fall_no = as_text(row.get("Fall") or row.get("Med. Fallnr.") or row.get("Medizinischer Fall"))
    fall_hash = stable_hash("fall", fall_no, salt=salt) if fall_no else None
    lookup_column = "Fallnummer" if include_identifiers and fall_no else "FallHash"
    lookup_value = fall_no if lookup_column == "Fallnummer" else fall_hash
    if lookup_value:
        found = con.execute(
            f"""
            SELECT AufenthaltID
            FROM tblAufenthalt
            WHERE GelenkFallID = ? AND EpisodeID = ? AND {lookup_column} = ?
            """,
            (gelenk_fall_id, episode_id, lookup_value),
        ).fetchone()
        if found:
            return int(found[0])

    cur = con.execute(
        """
        INSERT INTO tblAufenthalt (
            GelenkFallID, EpisodeID, Fallnummer, FallHash, MedizinischerFall, Station, Quelle
        )
        VALUES (?, ?, ?, ?, ?, ?, 'Orthodok')
        """,
        (
            gelenk_fall_id,
            episode_id,
            fall_no if include_identifiers else None,
            fall_hash,
            as_text(row.get("Medizinischer Fall")) if include_identifiers else None,
            as_text(row.get("Aufnehmende fachl.OE")),
        ),
    )
    return int(cur.lastrowid)


def get_or_create_intervention(
    con: sqlite3.Connection,
    stay_id: int,
    op_date: str | None,
    intervention_type_id: int,
    intervention_text: str,
) -> int:
    found = con.execute(
        """
        SELECT EingriffID
        FROM tblEingriff
        WHERE AufenthaltID = ? AND OPDatum IS ? AND EingriffsartID = ?
        """,
        (stay_id, op_date, intervention_type_id),
    ).fetchone()
    if found:
        return int(found[0])
    cur = con.execute(
        """
        INSERT INTO tblEingriff (AufenthaltID, OPDatum, EingriffsartID, EingriffsartTextQuelle, Quelle)
        VALUES (?, ?, ?, ?, 'Orthodok')
        """,
        (stay_id, op_date, intervention_type_id, intervention_text),
    )
    return int(cur.lastrowid)


def yesish(*values: Any) -> str | None:
    for value in values:
        text = as_text(value)
        if text:
            return text
    return None


def plausible_access_text(value: Any) -> str | None:
    text = as_text(value)
    if not text:
        return None
    compact = "".join(ch for ch in text if ch.lower() in "0123456789abcdef")
    if len(compact) >= 24 and len(compact) == len(text.replace("-", "").replace("_", "")):
        return None
    if text.lower() in {"x", "ja", "yes", "true", "1", "n.e", "n.e.", "n. e", "n. e.", "sonstige"}:
        return None
    return text


def orthodok_access(row: dict[str, Any]) -> str | None:
    access_fields = (
        ("Zugänge - medial parapatellar", "medial parapatellar"),
        ("Zugänge - lateral parapatellar", "lateral parapatellar"),
        ("Zugänge - Midvastus", "Midvastus"),
        ("Zugänge - Subvastus", "Subvastus"),
        ("Zugänge - Keblish", "Keblish/Keblisch"),
        ("Zugänge - Z Plastik", "Z-Plastik"),
        ("Zugaenge - medial parapatellar", "medial parapatellar"),
        ("Zugaenge - lateral parapatellar", "lateral parapatellar"),
        ("Zugaenge - Midvastus", "Midvastus"),
        ("Zugaenge - Subvastus", "Subvastus"),
        ("Zugaenge - Keblish", "Keblish/Keblisch"),
        ("Zugaenge - Z Plastik", "Z-Plastik"),
        ("Zugang - transgluteal/lateral", "transgluteal/lateral"),
        ("Zugang - transgluteal", "transgluteal/lateral"),
        ("Zugang - lateral", "lateral"),
    )
    terms: list[str] = []
    for header, label in access_fields:
        if as_text(row.get(header)):
            terms.append(label)
    for header in ("Zugänge - andere Text", "Zugaenge - andere Text", "Zugang - andere Text"):
        if text := plausible_access_text(row.get(header)):
            terms.append(text)
    unique_terms: list[str] = []
    for term in terms:
        if term not in unique_terms:
            unique_terms.append(term)
    return "; ".join(unique_terms) or None


def first_text(row: dict[str, Any], *headers: str) -> str | None:
    for header in headers:
        text = as_text(row.get(header))
        if text:
            return text
    return None


def insert_orthodok_op(
    con: sqlite3.Connection,
    *,
    intervention_id: int,
    raw_id: int,
    row: dict[str, Any],
) -> None:
    access = orthodok_access(row)
    navigation = yesish(
        row.get("Achsenbestimmung Femur - NAV"),
        row.get("Achsenbestimmung Femur - PSI"),
        row.get("Achsenbestimmung Femur - konventionell"),
    )
    fixation = infer_fixation(row)
    con.execute(
        """
        INSERT INTO tblOrthodokOP (
            EingriffID, RawOrthodokID, Version, ImplantatSystem, FemurGroesse,
            FemurTyp, FemurZementiert, TibiaGroesse, TibiaTyp, TibiaZementiert,
            PatellaErsatz, PatellaGroesse, PatellaZementiert, GleitflaecheGroesse,
            GleitflaecheTyp, Fixation, Zugang, Navigation, BlutsperreDauer,
            BlutsperreDruck, RohdatenJSON, Quelle
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'Orthodok')
        """,
        (
            intervention_id,
            raw_id,
            as_text(row.get("Version")),
            first_text(row, "System - Textfeld", "HEP Pfannentyp Bezeichnung", "HTEP OP-Befund - OrthoDok HTEP OP-Befund "),
            first_text(row, "Femur - Groesse", "Femur - Größe", "HEP Schaft Größe Bezeichner"),
            first_text(row, "Femur - Typ", "HEP Schafttyp - Bezeichner"),
            first_text(row, "Femur - zementiert", "Schaft zementiert", "HEP Schaft zementiert"),
            as_text(row.get("Tibia - Groesse") or row.get("Tibia - Größe")),
            as_text(row.get("Tibia - Typ")),
            as_text(row.get("Tibia - zementiert")),
            as_text(row.get("Patella - Ersatz")),
            as_text(row.get("Patella - Groesse") or row.get("Patella - Größe")),
            as_text(row.get("Patella - zementiert")),
            first_text(row, "Gleitflaeche - Groesse", "Gleitfläche - Größe", "HEP Pfannengröße Bezeichner", "HEP Schaft Kopfdurchmesser"),
            first_text(row, "Gleitflaeche - Typ", "Gleitfläche - Typ", "HEP Inlay - Größe Bezeichner", "HEP Pfanne Inlayform Bezeichner"),
            fixation,
            access,
            navigation,
            as_text(row.get("Ortho: Esmarchsche Blutsperre - Dauer")),
            as_text(row.get("Ortho: Esmarchsche Blutsperre - Druck")),
            json.dumps(row, ensure_ascii=False, default=str),
        ),
    )


def infer_fixation(row: dict[str, Any]) -> str | None:
    femur = first_text(row, "Femur - zementiert", "Schaft zementiert", "HEP Schaft zementiert")
    tibia = as_text(row.get("Tibia - zementiert"))
    patella = as_text(row.get("Patella - zementiert"))
    values = [v.lower() for v in [femur, tibia, patella] if v]
    if not values:
        return None
    zement_count = sum("ja" in v or v in {"1", "true", "x"} for v in values)
    if zement_count == len(values):
        return "zementiert"
    if zement_count == 0:
        return "zementfrei"
    return "hybrid"


def masked_raw(row: dict[str, Any], include_identifiers: bool) -> dict[str, Any]:
    if include_identifiers:
        return row
    result = dict(row)
    for key in DIRECT_IDENTIFIER_HEADERS:
        if key in result:
            result[key] = None
    return result


def import_orthodok(
    excel_path: Path,
    db_path: Path,
    *,
    sheet_name: str,
    limit: int | None,
    include_identifiers: bool,
    salt: str,
    joint: str,
) -> int:
    db_path.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(db_path) as con:
        con.execute("PRAGMA foreign_keys = ON")
        schema = (ROOT / "schema" / "pji_register_schema.sql").read_text(encoding="utf-8")
        con.executescript(schema)
        seed_antibiotics(con)
        migrate_episode_model(con)
        source_id = get_or_create_source(con, excel_path)
        batch_id = con.execute(
            """
            INSERT INTO tblImportBatch (SourceFileID, ImportType)
            VALUES (?, 'orthodok_excel')
            """,
            (source_id,),
        ).lastrowid

        wb = load_workbook(excel_path, read_only=True, data_only=True)
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet not found: {sheet_name}. Available: {', '.join(wb.sheetnames)}")
        ws = wb[sheet_name]
        iterator = ws.iter_rows(values_only=True)
        headers = normalized_headers(next(iterator))
        imported = 0
        try:
            for row_number, values in enumerate(iterator, start=2):
                if limit is not None and imported >= limit:
                    break
                raw = dict_from_row(headers, values)
                if not any(clean(v) is not None for v in raw.values()):
                    continue

                patient_id = get_or_create_patient(
                    con,
                    raw,
                    include_identifiers=include_identifiers,
                    salt=salt,
                )
                op_date = as_iso_date(raw.get("OP Datum"))
                intervention_type_id, intervention_text = infer_intervention_type(as_text(raw.get("Version")), raw)
                joint_id = get_or_create_joint_case(
                    con,
                    patient_id,
                    joint,
                    infer_side(raw.get("OP Seite")),
                    op_date,
                    intervention_type_id,
                )
                episode_id = get_or_create_episode(con, joint_id, op_date, intervention_type_id, intervention_text)
                stay_id = get_or_create_stay(
                    con,
                    joint_id,
                    episode_id,
                    raw,
                    include_identifiers=include_identifiers,
                    salt=salt,
                )
                intervention_id = get_or_create_intervention(
                    con,
                    stay_id,
                    op_date,
                    intervention_type_id,
                    intervention_text,
                )
                cur = con.execute(
                    """
                    INSERT INTO tblRawOrthodok
                        (ImportBatchID, SheetName, RowNumber, EingriffID, RawJSON)
                    VALUES (?, ?, ?, ?, ?)
                    """,
                    (
                        batch_id,
                        sheet_name,
                        row_number,
                        intervention_id,
                        json.dumps(masked_raw(raw, include_identifiers), ensure_ascii=False, default=str),
                    ),
                )
                insert_orthodok_op(con, intervention_id=intervention_id, raw_id=int(cur.lastrowid), row=raw)
                imported += 1
                if imported % 1000 == 0:
                    con.commit()

            con.execute(
                """
                UPDATE tblImportBatch
                SET FinishedAt = CURRENT_TIMESTAMP, RowCount = ?, Status = 'finished'
                WHERE ImportBatchID = ?
                """,
                (imported, batch_id),
            )
            con.commit()
            return imported
        except Exception:
            con.execute(
                """
                UPDATE tblImportBatch
                SET FinishedAt = CURRENT_TIMESTAMP, RowCount = ?, Status = 'failed'
                WHERE ImportBatchID = ?
                """,
                (imported, batch_id),
            )
            con.commit()
            raise


def main() -> None:
    parser = argparse.ArgumentParser(description="Import an OrthoDok Excel workbook into the PJI register.")
    parser.add_argument("excel_path", type=Path)
    parser.add_argument("--db", default=Path("outputs/pji_register.sqlite"), type=Path)
    parser.add_argument("--sheet", default=DEFAULT_SHEET)
    parser.add_argument("--joint", default="Knie", choices=["Knie", "Hüfte", "Huefte"])
    parser.add_argument("--limit", type=int)
    parser.add_argument("--include-identifiers", action="store_true")
    parser.add_argument(
        "--salt",
        default=os.environ.get("PJI_HASH_SALT", "pji-register-local-v1"),
        help="Salt for stable pseudonym hashes. Prefer setting PJI_HASH_SALT for real projects.",
    )
    args = parser.parse_args()
    count = import_orthodok(
        args.excel_path,
        args.db,
        sheet_name=args.sheet,
        limit=args.limit,
        include_identifiers=args.include_identifiers,
        salt=args.salt,
        joint="Hüfte" if args.joint == "Huefte" else args.joint,
    )
    print(f"Imported {count} OrthoDok rows into {args.db}")


if __name__ == "__main__":
    main()
