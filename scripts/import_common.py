from __future__ import annotations

import csv
import hashlib
import json
import sqlite3
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

try:
    from episode_model import migrate_episode_model
except ImportError:
    from scripts.episode_model import migrate_episode_model


NULL_MARKERS = {"", "#", "n.e.", "n.a.", "na", "none", "null", "-"}

DEFAULT_IDENTIFIER_COLUMNS = {
    "DOKUMENT",
    "FALL",
    "FALLNUMMER",
    "ORD_AUFNAHMENR",
    "PATIENT",
    "PAT_GEBOREN",
    "PAT_ORT",
    "PAT_PLZ",
    "PAT_STRASSE",
    "PAT_SVNR",
}
IDENTIFIER_NAME_PARTS = (
    "ADRESSE",
    "ADDRESS",
    "GEBURT",
    "GEBOREN",
    "NAME",
    "PLZ",
    "STRASSE",
    "SVNR",
)


def clean(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        if value.lower() in NULL_MARKERS:
            return None
    return value


def as_text(value: Any) -> str | None:
    value = clean(value)
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def as_iso_date(value: Any) -> str | None:
    value = clean(value)
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return text


def as_float(value: Any) -> float | None:
    value = clean(value)
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "")
    text = text.replace("<", "").replace(">", "")
    if "," in text and "." not in text:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def stable_hash(*parts: Any, salt: str) -> str:
    raw = "|".join("" if p is None else str(p) for p in parts)
    return hashlib.sha256(f"{salt}|{raw}".encode("utf-8")).hexdigest()[:32]


def is_identifier_column(header: str, extra_columns: set[str] | None = None) -> bool:
    normalized = header.upper().strip()
    if normalized in DEFAULT_IDENTIFIER_COLUMNS:
        return True
    if extra_columns and normalized in extra_columns:
        return True
    if normalized.startswith("PAT_"):
        return True
    return any(part in normalized for part in IDENTIFIER_NAME_PARTS)


def mask_identifier_row(
    row: dict[str, Any],
    *,
    include_identifiers: bool = False,
    extra_columns: list[str] | tuple[str, ...] | set[str] | None = None,
    salt: str = "pji-register-local-v1",
) -> dict[str, Any]:
    if include_identifiers:
        return dict(row)

    extra = {column.upper().strip() for column in extra_columns or []}
    masked: dict[str, Any] = {}
    for key, raw_value in row.items():
        if not is_identifier_column(key, extra):
            masked[key] = raw_value
            continue
        text = as_text(raw_value)
        masked[key] = f"[masked:{stable_hash(key, text, salt=salt)}]" if text else None
    return masked


def normalized_headers(values: list[Any] | tuple[Any, ...]) -> list[str]:
    seen: dict[str, int] = {}
    headers: list[str] = []
    for index, value in enumerate(values, start=1):
        header = as_text(value) or f"Unnamed_{index}"
        if header in seen:
            seen[header] += 1
            header = f"{header}_{seen[header]}"
        else:
            seen[header] = 1
        headers.append(header)
    return headers


def read_table(path: Path, sheet_name: str | None = None) -> tuple[str | None, list[dict[str, Any]]]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        wb = load_workbook(path, read_only=True, data_only=True)
        sheet = sheet_name or wb.sheetnames[0]
        ws = wb[sheet]
        rows = ws.iter_rows(values_only=True)
        headers = normalized_headers(next(rows))
        data = []
        for row in rows:
            mapped = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}
            if any(clean(v) is not None for v in mapped.values()):
                data.append(mapped)
        return sheet, data
    if suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else None
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            sample = f.read(4096)
            f.seek(0)
            if delimiter is None:
                delimiter = csv.Sniffer().sniff(sample, delimiters=";,|\t").delimiter
            reader = csv.DictReader(f, delimiter=delimiter)
            headers = normalized_headers(reader.fieldnames or [])
            data = []
            for raw in reader:
                mapped = {headers[i]: raw.get((reader.fieldnames or [])[i]) for i in range(len(headers))}
                if any(clean(v) is not None for v in mapped.values()):
                    data.append(mapped)
        return None, data
    raise ValueError(f"Unsupported table file type: {path.suffix}")


def connect_db(db_path: Path) -> sqlite3.Connection:
    con = sqlite3.connect(db_path)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys = ON")
    schema_path = Path(__file__).resolve().parents[1] / "schema" / "pji_register_schema.sql"
    con.executescript(schema_path.read_text(encoding="utf-8"))
    migrate_episode_model(con)
    con.commit()
    return con


def get_or_create_source(con: sqlite3.Connection, source_path: Path, source_system: str) -> int:
    digest = file_sha256(source_path)
    row = con.execute(
        "SELECT SourceFileID FROM tblSourceFile WHERE FileHashSHA256 = ? AND SourceSystem = ?",
        (digest, source_system),
    ).fetchone()
    if row:
        return int(row["SourceFileID"])
    cur = con.execute(
        """
        INSERT INTO tblSourceFile (SourceSystem, OriginalPath, FileName, FileHashSHA256)
        VALUES (?, ?, ?, ?)
        """,
        (source_system, str(source_path), source_path.name, digest),
    )
    return int(cur.lastrowid)


def start_import_batch(con: sqlite3.Connection, source_file_id: int, import_type: str) -> int:
    cur = con.execute(
        "INSERT INTO tblImportBatch (SourceFileID, ImportType) VALUES (?, ?)",
        (source_file_id, import_type),
    )
    return int(cur.lastrowid)


def finish_import_batch(con: sqlite3.Connection, batch_id: int, row_count: int, status: str = "finished", note: str | None = None) -> None:
    con.execute(
        """
        UPDATE tblImportBatch
        SET FinishedAt = CURRENT_TIMESTAMP, RowCount = ?, Status = ?, Note = ?
        WHERE ImportBatchID = ?
        """,
        (row_count, status, note, batch_id),
    )


def resolve_episode_context(
    con: sqlite3.Connection,
    *,
    episode_id: int | None = None,
    eingriff_id: int | None = None,
    fallnummer: str | None = None,
    patient_no: str | None = None,
    salt: str = "pji-register-local-v1",
) -> dict[str, Any]:
    if episode_id:
        row = con.execute(
            """
            SELECT ep.EpisodeID, NULL AS AufenthaltID, NULL AS EingriffID
            FROM tblEpisode ep
            WHERE ep.EpisodeID = ?
            """,
            (episode_id,),
        ).fetchone()
        if row:
            return dict(row)
    if eingriff_id:
        row = con.execute(
            """
            SELECT a.EpisodeID, a.AufenthaltID, e.EingriffID
            FROM tblEingriff e
            JOIN tblAufenthalt a ON a.AufenthaltID = e.AufenthaltID
            WHERE e.EingriffID = ?
            """,
            (eingriff_id,),
        ).fetchone()
        if row:
            return dict(row)
    fall_hash = stable_hash("fall", fallnummer, salt=salt) if fallnummer else None
    if fallnummer or fall_hash:
        row = con.execute(
            """
            SELECT a.EpisodeID, a.AufenthaltID, NULL AS EingriffID
            FROM tblAufenthalt a
            WHERE (a.Fallnummer = ? AND ? IS NOT NULL)
               OR (a.FallHash = ? AND ? IS NOT NULL)
            ORDER BY a.AufenthaltID DESC
            LIMIT 1
            """,
            (fallnummer, fallnummer, fall_hash, fall_hash),
        ).fetchone()
        if row:
            result = dict(row)
            result["FallHash"] = fall_hash
            return result
    patient_hash = stable_hash(f"patient:{patient_no}", salt=salt) if patient_no else None
    if patient_hash:
        row = con.execute(
            """
            SELECT ep.EpisodeID, NULL AS AufenthaltID, NULL AS EingriffID
            FROM tblEpisode ep
            JOIN tblGelenkFall g ON g.GelenkFallID = ep.GelenkFallID
            JOIN tblPatient p ON p.PatientID = g.PatientID
            WHERE p.PatientHash = ?
            ORDER BY ep.EpisodeBeginn IS NULL, ep.EpisodeBeginn DESC, ep.EpisodeID DESC
            LIMIT 1
            """,
            (patient_hash,),
        ).fetchone()
        if row:
            result = dict(row)
            result["PatientHash"] = patient_hash
            return result
    return {
        "EpisodeID": None,
        "AufenthaltID": None,
        "EingriffID": None,
        "FallHash": fall_hash,
        "PatientHash": patient_hash,
    }


def intervention_link_decision(
    con: sqlite3.Connection,
    episode_id: int | None,
    target_date: str | None,
    *,
    max_days: int = 30,
) -> dict[str, Any]:
    """Return a conservative, explainable intervention link decision.

    A microbiology or document row must not be attached to an arbitrarily old
    intervention merely because it is the only intervention in an episode.
    Exact/near dates may be linked automatically; distant or ambiguous dates
    remain episode-level review items.
    """
    if not episode_id:
        return {"EingriffID": None, "AbstandTage": None, "Score": 0.0, "Grund": "keine_episode"}
    if not target_date:
        return {"EingriffID": None, "AbstandTage": None, "Score": 0.0, "Grund": "datum_fehlt"}

    rows = con.execute(
        """
        SELECT e.EingriffID, e.OPDatum,
               CAST(ABS(julianday(e.OPDatum) - julianday(?)) AS INTEGER) AS AbstandTage
        FROM tblEingriff e
        JOIN tblAufenthalt a ON a.AufenthaltID = e.AufenthaltID
        WHERE a.EpisodeID = ?
          AND e.OPDatum IS NOT NULL
        ORDER BY AbstandTage, e.EingriffID
        """,
        (target_date, episode_id),
    ).fetchall()
    if not rows:
        return {"EingriffID": None, "AbstandTage": None, "Score": 0.0, "Grund": "kein_datiertes_verfahren"}

    best = rows[0]
    distance = int(best["AbstandTage"])
    if distance > max_days:
        return {"EingriffID": None, "AbstandTage": distance, "Score": 0.0, "Grund": "zeitlich_fern"}
    if len(rows) > 1 and int(rows[1]["AbstandTage"]) == distance:
        return {"EingriffID": None, "AbstandTage": distance, "Score": 0.0, "Grund": "mehrdeutig"}

    score = 1.0 if distance == 0 else max(0.55, round(1.0 - (distance / (max_days + 1)), 2))
    return {
        "EingriffID": int(best["EingriffID"]),
        "AbstandTage": distance,
        "Score": score,
        "Grund": "gleiches_datum" if distance == 0 else "zeitlich_nah",
    }


def nearest_intervention_for_episode(
    con: sqlite3.Connection,
    episode_id: int | None,
    target_date: str | None,
    *,
    max_days: int = 30,
) -> int | None:
    """Compatibility wrapper using the conservative linking rule."""
    return intervention_link_decision(
        con,
        episode_id,
        target_date,
        max_days=max_days,
    )["EingriffID"]


def row_to_json(row: dict[str, Any]) -> str:
    return json.dumps(row, ensure_ascii=False, default=str)
