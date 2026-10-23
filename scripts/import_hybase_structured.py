#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
import logging
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

try:
    from import_common import (
        as_float,
        as_iso_date,
        as_text,
        clean,
        connect_db,
        file_sha256,
        finish_import_batch,
        get_or_create_source,
        mask_identifier_row,
        normalized_headers,
        row_to_json,
        stable_hash,
        start_import_batch,
    )
except ImportError:
    from scripts.import_common import (
        as_float,
        as_iso_date,
        as_text,
        clean,
        connect_db,
        file_sha256,
        finish_import_batch,
        get_or_create_source,
        mask_identifier_row,
        normalized_headers,
        row_to_json,
        stable_hash,
        start_import_batch,
    )


LOGGER = logging.getLogger("hybase_structured")

EMPTY_RESULTS = {"", "#", "-", "n.e.", "n.a.", "na", "none", "null", "nicht getestet"}
SHEETS_EXCLUDED_FROM_DEFAULT = {"history"}
GERMAN_UMLAUTS = str.maketrans({"ä": "ae", "ö": "oe", "ü": "ue", "ß": "ss"})
THERAPY_HINT_WORDS = (
    "amoxicillin",
    "ampicillin",
    "cefazolin",
    "cefuroxim",
    "ciprofloxacin",
    "clindamycin",
    "daptomycin",
    "fosfomycin",
    "gentamicin",
    "linezolid",
    "meropenem",
    "metronidazol",
    "piperacillin",
    "rifampicin",
    "teicoplanin",
    "vancomycin",
)


PATIENT_FIELDS = [
    "PatientID",
    "HYBASEPatientID_Raw",
    "PatientHash",
    "Geburtsdatum",
    "Geburtsjahr",
    "Geschlecht",
    "Quelle",
]

LABORAUFTRAG_FIELDS = [
    "LaborauftragID",
    "SourceKey",
    "PatientID",
    "PatientHash",
    "LaborNr_Raw",
    "Auftragsdatum",
    "Befunddatum",
    "BefundlaufzeitTage",
    "Material_Raw",
    "Lokalisation_Raw",
    "Info_Raw",
    "KlinischeRelevanz_Raw",
    "Status_Raw",
    "Klassifizierung_Raw",
    "Fallnummer_Raw",
    "Station_Raw",
    "Fachbereich_Raw",
    "Freitext_Raw",
    "AntibiotikaHinweis_Raw",
    "Flag_AntibiotikaImFreitext",
    "Quellblatt",
    "Quelldatei",
    "RawJSON",
    "Flag_BefunddatumVorAuftragsdatum",
    "Flag_BefundlaufzeitGroesser7Tage",
    "Flag_BefundlaufzeitGroesser14Tage",
    "Flag_FehlendesAuftragsdatum",
    "Flag_FehlendesBefunddatum",
]

PROBE_FIELDS = [
    "ProbeID",
    "SourceKey",
    "LaborauftragID",
    "LaborauftragSourceKey",
    "Material_Raw",
    "Material_Standard",
    "Probenbeschreibung_Raw",
    "Lokalisation_Raw",
    "Gelenk_Abgeleitet",
    "Seite_Abgeleitet",
    "AnatomischeStruktur",
    "Implantatbezug_Abgeleitet",
    "ImplantattypHinweis",
    "Gelenk_Ableitungsquelle",
    "Seite_Ableitungsquelle",
    "Flag_GelenkUnsicher",
    "Flag_SeiteUnsicher",
    "PJI_Relevanz_Vorlaeufig",
    "AnzahlErreger",
    "AnzahlKulturellPositiveErreger",
    "Polymikrobiell",
    "NurNiedrigVirulenteErreger",
    "MindestensEinHochvirulenterErreger",
    "MindestensEinResistenterErreger",
    "AntibiogrammVorhanden",
    "NurAnreicherungsnachweis",
    "GelenkFallID",
    "EpisodeID",
    "AufenthaltID",
    "EingriffID",
    "OPDatum",
    "TageDifferenzZuOP",
    "Zuordnungssicherheit",
]

ISOLAT_FIELDS = [
    "IsolatID",
    "SourceKey",
    "ProbeID",
    "ProbeSourceKey",
    "LaborauftragID",
    "LaborauftragSourceKey",
    "Erreger_Raw",
    "Erreger_Standard",
    "Erregergruppe",
    "Gramstatus",
    "Virulenzgruppe",
    "Resistenzmarker_Raw",
    "ResistenzmarkerDatum",
    "Subtyp_Raw",
    "Ergebnisstatus_Raw",
    "Messwert_Raw",
    "Kommentar_Raw",
    "Positiv",
    "Negativ",
    "ErgebnisUnklar",
    "KeinMikrobiologischerErreger",
    "Nachweismethode",
    "NurAnreicherung",
    "NachBlutkulturflasche",
    "LangsamWachsenderErreger",
    "KeimmengeKategorie",
    "SonikationKBEKategorie",
    "QuantitativerWert",
    "QuantitativeEinheit",
]

ANTIBIOGRAMM_FIELDS = [
    "AntibiogrammID",
    "SourceKey",
    "IsolatID",
    "IsolatSourceKey",
    "Antibiotikum_Raw",
    "Antibiotikum_Standard",
    "Antibiotikaklasse",
    "Spezialtest",
    "Ergebnis_Raw",
    "SIR_Standard",
    "MIC_Raw",
    "Methode_Raw",
    "Breakpointsystem",
    "Testdatum",
    "AntibiotikumSpalte",
]

RAW_LINK_FIELDS = [
    "SheetName",
    "RowNumber",
    "RawJSON",
    "LaborauftragSourceKey",
    "ProbeSourceKey",
    "IsolatSourceKey",
    "AntibiogrammSourceKeys",
]


def normalized_key(value: object | None) -> str:
    text = as_text(value) or ""
    text = text.lower().translate(GERMAN_UMLAUTS)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", "", text)


def raw_text(value: object | None) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def date_to_year(value: object | None) -> int | None:
    text = as_iso_date(value)
    if not text or len(text) < 4:
        return None
    try:
        return int(text[:4])
    except ValueError:
        return None


def parse_date(value: str | None) -> date | None:
    if not value:
        return None
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError:
        return None


def day_diff(start: str | None, end: str | None) -> int | None:
    start_date = parse_date(start)
    end_date = parse_date(end)
    if not start_date or not end_date:
        return None
    return (end_date - start_date).days


def load_config(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def resolve_path(path: str | None, *, base_dir: Path) -> Path | None:
    if not path:
        return None
    candidate = Path(path)
    return candidate if candidate.is_absolute() else base_dir / candidate


def load_mapping(path: Path, raw_column: str) -> tuple[dict[str, dict[str, str]], list[dict[str, str]]]:
    rows: list[dict[str, str]] = []
    lookup: dict[str, dict[str, str]] = {}
    if not path.exists():
        return lookup, rows
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            rows.append(dict(row))
            if str(row.get("Aktiv", "1")).strip() in {"0", "false", "False", "nein"}:
                continue
            raw = row.get(raw_column)
            if raw:
                lookup[normalized_key(raw)] = dict(row)
    return lookup, rows


def first_value(row: dict[str, Any], candidates: list[str]) -> Any:
    for column in candidates:
        if column in row and clean(row[column]) is not None:
            return row[column]
    return None


def col(config: dict[str, Any], name: str) -> list[str]:
    return list(config["columns"].get(name, []))


def is_ab_column(header: str, prefixes: list[str]) -> bool:
    return any(header.startswith(prefix) for prefix in prefixes)


def antibiotic_from_header(header: str) -> str:
    name = header[3:] if header.startswith("AB_") else header
    name = re.sub(r"_\d+$", "", name)
    return name.strip()


def standardize_material(
    material_raw: str | None,
    mapping: dict[str, dict[str, str]],
) -> tuple[str, bool]:
    key = normalized_key(material_raw)
    if key in mapping:
        return mapping[key].get("Material_Standard") or "UNKLAR", True
    text = key
    if not text:
        return "UNKLAR", False
    if any(term in text for term in ("implantat", "prothese", "sonikation", "soni")):
        return "SONIKATION_IMPLANTAT", False
    if any(term in text for term in ("gewebe", "synovia", "synovialis", "kapsel", "membran")):
        return "GEWEBE", False
    if "knochen" in text:
        return "KNOCHEN", False
    if "punktat" in text and any(term in text for term in ("gelenk", "hueft", "knie", "schulter")):
        return "GELENKPUNKTAT", False
    if "punktat" in text:
        return "SONSTIGES_PUNKTAT", False
    if "abstrich" in text:
        return "ABSTRICH", False
    if "wundsekret" in text or ("sekret" in text and "wund" in text):
        return "WUNDSEKRET", False
    if "haematom" in text:
        return "HAEMATOM", False
    if "abszess" in text:
        return "ABSZESS", False
    if "blut" in text or text.startswith("blk") or "blutkultur" in text:
        return "BLUTKULTUR", False
    if any(term in text for term in ("harn", "stuhl", "venenkatheter", "katheter", "zvk", "venflon")):
        return "SONSTIGES", False
    return "SONSTIGES", False


def derive_anatomy(*values: str | None) -> dict[str, Any]:
    raw = " ".join(value for value in values if value)
    text = normalized_key(raw)
    joint_patterns = [
        ("HUEFTE", ("huefte", "hueft", "cox", "pfanne", "schaft")),
        ("KNIE", ("knie", "tibia", "femurkomponente", "patella")),
        ("SCHULTER", ("schulter", "humerus")),
        ("ELLENBOGEN", ("ellenbogen", "ellbogen")),
        ("SPRUNGGELENK", ("sprunggelenk", "osg")),
        ("HAND", ("hand", "finger")),
        ("WIRBELSAEULE", ("wirbel", "spine", "spondyl", "lws", "bws", "hws")),
    ]
    found_joints = [joint for joint, terms in joint_patterns if any(term in text for term in terms)]
    if len(found_joints) == 1:
        joint = found_joints[0]
        joint_uncertain = 0
    elif len(found_joints) > 1:
        joint = "UNKLAR"
        joint_uncertain = 1
    else:
        joint = "UNKLAR"
        joint_uncertain = 1

    has_right = bool(re.search(r"(^|[^a-z])(re|rechts|right)([^a-z]|$)", raw.lower()))
    has_left = bool(re.search(r"(^|[^a-z])(li|links|left)([^a-z]|$)", raw.lower()))
    if has_right and has_left:
        side = "BEIDSEITS"
        side_uncertain = 0
    elif has_right:
        side = "RECHTS"
        side_uncertain = 0
    elif has_left:
        side = "LINKS"
        side_uncertain = 0
    else:
        side = "UNKLAR"
        side_uncertain = 1

    implant_hint = "JA" if any(term in text for term in ("implantat", "prothese", "pfanne", "schaft", "inlay")) else "UNKLAR"
    return {
        "Gelenk_Abgeleitet": joint,
        "Seite_Abgeleitet": side,
        "AnatomischeStruktur": raw or None,
        "Implantatbezug_Abgeleitet": implant_hint,
        "ImplantattypHinweis": raw if implant_hint == "JA" else None,
        "Gelenk_Ableitungsquelle": "Material/Lokalisation/Freitext" if raw else None,
        "Seite_Ableitungsquelle": "Material/Lokalisation/Freitext" if raw else None,
        "Flag_GelenkUnsicher": joint_uncertain,
        "Flag_SeiteUnsicher": side_uncertain,
    }


def pji_relevance(material: str, joint: str, implant_hint: str, *context_values: str | None) -> str:
    context = normalized_key(" ".join(value for value in context_values if value))
    non_pji_terms = (
        "nase",
        "rachen",
        "harn",
        "stuhl",
        "venenkatheter",
        "zvk",
        "venflon",
        "arteriell",
        "haende",
        "handeflaeche",
        "abklatsch",
        "hautflaeche",
    )
    if any(term in context for term in non_pji_terms):
        return "NICHT_PJI_RELEVANT"
    if material == "BLUTKULTUR" and joint == "UNKLAR":
        return "NICHT_PJI_RELEVANT"
    if material in {"GEWEBE", "KNOCHEN", "GELENKPUNKTAT", "SONIKATION_IMPLANTAT"} and joint in {"HUEFTE", "KNIE"}:
        return "PJI_RELEVANT_WAHRSCHEINLICH"
    if material == "SONIKATION_IMPLANTAT" and implant_hint == "JA":
        return "PJI_RELEVANT_MOEGLICH"
    if material in {"ABSTRICH", "WUNDSEKRET", "HAEMATOM", "ABSZESS"} and joint in {"HUEFTE", "KNIE"}:
        return "PJI_RELEVANT_MOEGLICH"
    if material in {"BLUTKULTUR", "SONSTIGES"} and joint in {"UNKLAR", "WIRBELSAEULE", "HAND"}:
        return "NICHT_PJI_RELEVANT"
    return "UNKLAR"


def organism_fallback(raw: str | None) -> dict[str, str]:
    text = normalized_key(raw)
    if not text or any(term in text for term in ("steril", "keinerreger", "keinwachstum", "negativ")):
        return {
            "Erreger_Standard": "steril/kein Erreger",
            "Erregergruppe": "NICHT_ERREGERBEFUND",
            "Gramstatus": "KEIN_ERREGER",
            "Virulenzgruppe": "NIEDRIGVIRULENT",
        }
    if "propionibacteriumacnes" in text or "cutibacteriumacnes" in text:
        return {
            "Erreger_Standard": "Cutibacterium acnes",
            "Erregergruppe": "CUTIBACTERIUM",
            "Gramstatus": "GRAM_POSITIV",
            "Virulenzgruppe": "NIEDRIGVIRULENT",
        }
    if "staph" in text or "staphylococcus" in text:
        if "aureus" in text:
            return {
                "Erreger_Standard": "Staphylococcus aureus",
                "Erregergruppe": "STAPHYLOCOCCUS_AUREUS",
                "Gramstatus": "GRAM_POSITIV",
                "Virulenzgruppe": "HOCHVIRULENT",
            }
        return {
            "Erreger_Standard": raw or "KoNS nicht weiter differenziert",
            "Erregergruppe": "KONS",
            "Gramstatus": "GRAM_POSITIV",
            "Virulenzgruppe": "NIEDRIGVIRULENT",
        }
    if "strep" in text or "streptococcus" in text:
        return {
            "Erreger_Standard": raw or "Streptococcus spp.",
            "Erregergruppe": "STREPTOCOCCUS",
            "Gramstatus": "GRAM_POSITIV",
            "Virulenzgruppe": "UNKLAR",
        }
    if "enterococcus" in text:
        return {
            "Erreger_Standard": raw or "Enterococcus spp.",
            "Erregergruppe": "ENTEROCOCCUS",
            "Gramstatus": "GRAM_POSITIV",
            "Virulenzgruppe": "UNKLAR",
        }
    if "escherichiacoli" in text or text == "ecoli":
        return {
            "Erreger_Standard": "Escherichia coli",
            "Erregergruppe": "ENTEROBACTERALES",
            "Gramstatus": "GRAM_NEGATIV",
            "Virulenzgruppe": "HOCHVIRULENT",
        }
    if "pseudomonas" in text:
        return {
            "Erreger_Standard": raw or "Pseudomonas spp.",
            "Erregergruppe": "PSEUDOMONAS",
            "Gramstatus": "GRAM_NEGATIV",
            "Virulenzgruppe": "HOCHVIRULENT",
        }
    if "coryne" in text:
        return {
            "Erreger_Standard": raw or "Corynebacterium spp.",
            "Erregergruppe": "CORYNEBACTERIUM",
            "Gramstatus": "GRAM_POSITIV",
            "Virulenzgruppe": "NIEDRIGVIRULENT",
        }
    if any(term in text for term in ("candida", "aspergillus", "sprosspilz")):
        return {
            "Erreger_Standard": raw or "Pilz",
            "Erregergruppe": "PILZE",
            "Gramstatus": "PILZ",
            "Virulenzgruppe": "UNKLAR",
        }
    if any(term in text for term in ("anaerob", "bacteroides", "finegoldia", "micromonas", "parvimonas")):
        return {
            "Erreger_Standard": raw or "Anaerobier",
            "Erregergruppe": "ANAEROBIER",
            "Gramstatus": "ANAEROB",
            "Virulenzgruppe": "UNKLAR",
        }
    return {
        "Erreger_Standard": raw or "unklar",
        "Erregergruppe": "UNKLAR",
        "Gramstatus": "UNKLAR",
        "Virulenzgruppe": "UNKLAR",
    }


def standardize_organism(raw: str | None, mapping: dict[str, dict[str, str]]) -> tuple[dict[str, str], bool]:
    key = normalized_key(raw)
    if key in mapping:
        row = mapping[key]
        return (
            {
                "Erreger_Standard": row.get("Erreger_Standard") or raw or "unklar",
                "Erregergruppe": row.get("Erregergruppe") or "UNKLAR",
                "Gramstatus": row.get("Gramstatus") or "UNKLAR",
                "Virulenzgruppe": row.get("Virulenzgruppe") or "UNKLAR",
            },
            True,
        )
    return organism_fallback(raw), False


def infer_method(material_standard: str, *values: str | None) -> str:
    text = normalized_key(" ".join(value for value in values if value))
    if any(term in text for term in ("biofire", "pcr", "panel")):
        return "PCR_PANEL"
    if "screen" in text or "mrsa" in text and "screen" in text:
        return "SCREENING"
    if material_standard == "BLUTKULTUR" or "blutkultur" in text:
        return "BLUTKULTURFLASCHE"
    if material_standard == "SONIKATION_IMPLANTAT" or "sonikation" in text:
        return "SONIKATIONSKULTUR"
    if "anreicherung" in text:
        return "ANREICHERUNG"
    return "KULTUR"


def infer_result_flags(organism: dict[str, str], method: str, *values: str | None) -> dict[str, int]:
    text = normalized_key(" ".join(value for value in values if value))
    is_non_pathogen = organism.get("Erregergruppe") == "NICHT_ERREGERBEFUND"
    negative_text = any(term in text for term in ("steril", "keinerreger", "keinwachstum", "negativ"))
    positive_text = any(term in text for term in ("positiv", "pos", "wachstum", "nachweis"))
    if is_non_pathogen or negative_text:
        return {"Positiv": 0, "Negativ": 1, "ErgebnisUnklar": 0, "KeinMikrobiologischerErreger": 1}
    if organism.get("Erreger_Raw") or organism.get("Erreger_Standard") not in {None, "", "unklar"}:
        return {"Positiv": 1, "Negativ": 0, "ErgebnisUnklar": 0, "KeinMikrobiologischerErreger": 0}
    if positive_text and method in {"PCR_PANEL", "SCREENING"}:
        return {"Positiv": 1, "Negativ": 0, "ErgebnisUnklar": 0, "KeinMikrobiologischerErreger": 0}
    return {"Positiv": 0, "Negativ": 0, "ErgebnisUnklar": 1, "KeinMikrobiologischerErreger": 0}


def quantity_fields(*values: str | None) -> dict[str, Any]:
    raw = " ".join(value for value in values if value)
    text = normalized_key(raw)
    quantitative = re.search(r"([<>]?\s*\d+(?:[,.]\d+)?)\s*(kbe|cfu)", raw.lower())
    quant_value = None
    quant_unit = None
    if quantitative:
        quant_value = as_float(quantitative.group(1))
        quant_unit = quantitative.group(2).upper()
    if "anreicherung" in text:
        category = "NUR_ANREICHERUNG"
    elif any(term in text for term in ("einzelnbismaessig", "vereinzeltbismaessig")):
        category = "EINZELN_BIS_MAESSIG"
    elif any(term in text for term in ("einzeln", "vereinzelt")):
        category = "EINZELN"
    elif "maessig" in text:
        category = "MAESSIG"
    elif "reichlich" in text:
        category = "REICHLICH"
    elif "<50" in raw.replace(" ", "") and quantitative:
        category = "SONIKATION_UNTER_50_KBE_ML"
    elif quant_value is not None and quant_value >= 50:
        category = "SONIKATION_AB_50_KBE_ML"
    elif quantitative:
        category = "QUANTITATIV_SONSTIG"
    else:
        category = "NICHT_ANGEGEBEN"
    return {
        "NurAnreicherung": 1 if category == "NUR_ANREICHERUNG" else 0,
        "NachBlutkulturflasche": 1 if "blutkulturflasche" in text else 0,
        "LangsamWachsenderErreger": 1 if any(term in text for term in ("langsamwachsend", "spaetwachsend")) else 0,
        "KeimmengeKategorie": category,
        "SonikationKBEKategorie": category if category.startswith("SONIKATION") else None,
        "QuantitativerWert": quant_value,
        "QuantitativeEinheit": quant_unit,
    }


def antibiotic_hint(*values: str | None) -> tuple[str | None, int]:
    raw = " ".join(value for value in values if value)
    key = normalized_key(raw)
    hits = sorted({word for word in THERAPY_HINT_WORDS if normalized_key(word) in key})
    return ("; ".join(hits) if hits else None, 1 if hits else 0)


def normalize_sir(raw: str | None) -> tuple[str, str | None]:
    text = (raw or "").strip()
    key = normalized_key(text)
    if not key or key in {normalized_key(value) for value in EMPTY_RESULTS}:
        return "NICHT_GETESTET", None
    mic_raw = text if re.search(r"[<>]?\s*\d+[,.]?\d*", text) else None
    if key in {"s", "sensibel", "sensible", "susceptible"}:
        return "S", mic_raw
    if key in {"i", "intermediaer", "intermediar", "intermediate"}:
        return "I", mic_raw
    if key in {"r", "resistent", "resistant"}:
        return "R", mic_raw
    if text[:1].upper() in {"S", "I", "R"} and len(text) <= 4:
        return text[:1].upper(), mic_raw
    return "UNKLAR", mic_raw


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def copy_mapping_csv(path: Path, rows: list[dict[str, str]], fieldnames: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def make_quality_report(path: Path, summary: list[tuple[str, Any]], tables: dict[str, list[dict[str, Any]]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Kennzahl", "Wert"])
    for key, value in summary:
        ws.append([key, value])
    for sheet_name, rows in tables.items():
        ws = wb.create_sheet(sheet_name[:31])
        if not rows:
            ws.append(["keine Eintraege"])
            continue
        fieldnames = list(rows[0].keys())
        ws.append(fieldnames)
        for row in rows[:5000]:
            ws.append([row.get(field) for field in fieldnames])
    wb.save(path)


def write_database_import(path: Path, tables: dict[str, tuple[list[dict[str, Any]], list[str]]]) -> None:
    wb = Workbook()
    first = True
    for sheet_name, (rows, fields) in tables.items():
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = sheet_name[:31]
        ws.append(fields)
        for row in rows:
            ws.append([row.get(field) for field in fields])
    wb.save(path)


def process_hybase(args: argparse.Namespace) -> dict[str, Any]:
    config = load_config(args.config)
    base_dir = Path.cwd()
    mapping_paths = config.get("mapping_paths", {})
    material_map, material_map_rows = load_mapping(
        resolve_path(mapping_paths.get("material"), base_dir=base_dir) or Path(""),
        "Material_Raw",
    )
    organism_map, organism_map_rows = load_mapping(
        resolve_path(mapping_paths.get("organism"), base_dir=base_dir) or Path(""),
        "Erreger_Raw",
    )
    antibiotic_map, antibiotic_map_rows = load_mapping(
        resolve_path(mapping_paths.get("antibiotic"), base_dir=base_dir) or Path(""),
        "Antibiotikum_Raw",
    )

    wb = load_workbook(args.file, read_only=True, data_only=True)
    if args.sheets:
        sheets = [sheet.strip() for sheet in args.sheets.split(",") if sheet.strip()]
    else:
        sheets = list(config.get("sheets") or [config.get("default_sheet") or wb.sheetnames[0]])
    sheets = [sheet for sheet in sheets if sheet in wb.sheetnames and normalized_key(sheet) not in SHEETS_EXCLUDED_FROM_DEFAULT]
    if not sheets:
        raise ValueError("Keine passenden HYBASE-Tabellenblaetter gefunden.")

    source_digest = file_sha256(args.file)[:16]
    patients: dict[str, dict[str, Any]] = {}
    labor_orders: dict[str, dict[str, Any]] = {}
    probes: dict[str, dict[str, Any]] = {}
    isolates: dict[str, dict[str, Any]] = {}
    antibiograms: dict[str, dict[str, Any]] = {}
    raw_links: list[dict[str, Any]] = []
    isolate_by_probe: dict[str, set[str]] = defaultdict(set)
    antibiogram_by_isolate: dict[str, set[str]] = defaultdict(set)

    counters: Counter[str] = Counter()
    material_counts: Counter[str] = Counter()
    unmapped_materials: Counter[str] = Counter()
    unmapped_organisms: Counter[str] = Counter()
    unmapped_antibiotics: Counter[str] = Counter()
    invalid_sir: Counter[str] = Counter()
    antibiotic_columns_seen: Counter[str] = Counter()
    duplicate_raw_counter: Counter[str] = Counter()
    lab_patient_assignments: dict[str, set[str]] = defaultdict(set)

    patient_next = 1
    labor_next = 1
    probe_next = 1
    isolate_next = 1
    antibiogram_next = 1

    for sheet_name in sheets:
        LOGGER.info("Lese HYBASE-Blatt %s", sheet_name)
        ws = wb[sheet_name]
        rows = ws.iter_rows(values_only=True)
        try:
            headers = normalized_headers(next(rows))
        except StopIteration:
            continue
        ab_columns = [header for header in headers if is_ab_column(header, list(config.get("antibiogram_prefixes", ["AB_"])))]
        for excel_row_number, row_values in enumerate(rows, start=2):
            row = {headers[i]: row_values[i] if i < len(row_values) else None for i in range(len(headers))}
            if not any(clean(value) is not None for value in row.values()):
                continue
            if args.limit and counters["raw_rows"] >= args.limit:
                break
            counters["raw_rows"] += 1

            raw_signature = row_to_json({key: raw_text(value) for key, value in row.items() if clean(value) is not None})
            duplicate_raw_counter[raw_signature] += 1

            patient_raw = raw_text(first_value(row, col(config, "patient_raw")))
            patient_hash = stable_hash(f"patient:{patient_raw}", salt=args.salt) if patient_raw else stable_hash(
                "hybase_missing_patient", sheet_name, excel_row_number, salt=args.salt
            )
            if patient_hash not in patients:
                patients[patient_hash] = {
                    "PatientID": patient_next,
                    "HYBASEPatientID_Raw": patient_raw,
                    "PatientHash": patient_hash,
                    "Geburtsdatum": as_iso_date(first_value(row, col(config, "birth_date"))),
                    "Geburtsjahr": date_to_year(first_value(row, col(config, "birth_date"))),
                    "Geschlecht": raw_text(first_value(row, col(config, "sex"))),
                    "Quelle": "HYBASE",
                }
                patient_next += 1

            material_raw = raw_text(first_value(row, col(config, "material")))
            localization_raw = raw_text(first_value(row, col(config, "localization")))
            info_raw = raw_text(first_value(row, col(config, "info")))
            status_raw = raw_text(first_value(row, col(config, "result_status")))
            comment_raw = raw_text(first_value(row, col(config, "result_comment")))
            pathogen_raw = raw_text(first_value(row, col(config, "organism")))
            result_value_raw = raw_text(first_value(row, col(config, "result_value")))
            subtype_raw = raw_text(first_value(row, col(config, "subtype")))
            resistance_marker_raw = raw_text(first_value(row, col(config, "resistance_marker")))
            order_date = as_iso_date(first_value(row, col(config, "order_date")))
            report_date = as_iso_date(first_value(row, col(config, "report_date")))
            runtime_days = day_diff(order_date, report_date)
            lab_raw = raw_text(first_value(row, col(config, "labor_number")))
            if not lab_raw:
                counters["missing_labor_number"] += 1
            labor_source_key = (
                f"hybase:lab:{lab_raw}"
                if lab_raw
                else f"hybase:lab_missing:{source_digest}:{sheet_name}:{excel_row_number}"
            )
            lab_patient_assignments[labor_source_key].add(patient_hash)
            ab_hint, ab_hint_flag = antibiotic_hint(info_raw, comment_raw)
            if labor_source_key not in labor_orders:
                labor_orders[labor_source_key] = {
                    "LaborauftragID": labor_next,
                    "SourceKey": labor_source_key,
                    "PatientID": patients[patient_hash]["PatientID"],
                    "PatientHash": patient_hash,
                    "LaborNr_Raw": lab_raw,
                    "Auftragsdatum": order_date,
                    "Befunddatum": report_date,
                    "BefundlaufzeitTage": runtime_days,
                    "Material_Raw": material_raw,
                    "Lokalisation_Raw": localization_raw,
                    "Info_Raw": info_raw,
                    "KlinischeRelevanz_Raw": raw_text(first_value(row, col(config, "clinical_relevance"))),
                    "Status_Raw": status_raw,
                    "Klassifizierung_Raw": raw_text(first_value(row, col(config, "classification"))),
                    "Fallnummer_Raw": raw_text(first_value(row, col(config, "case_number"))),
                    "Station_Raw": raw_text(first_value(row, col(config, "station"))),
                    "Fachbereich_Raw": raw_text(first_value(row, col(config, "department"))),
                    "Freitext_Raw": " | ".join(value for value in (info_raw, status_raw, comment_raw) if value) or None,
                    "AntibiotikaHinweis_Raw": ab_hint,
                    "Flag_AntibiotikaImFreitext": ab_hint_flag,
                    "Quellblatt": sheet_name,
                    "Quelldatei": args.file.name,
                    "RawJSON": row_to_json(mask_identifier_row(row, include_identifiers=args.include_identifiers, salt=args.salt)),
                    "Flag_BefunddatumVorAuftragsdatum": 1 if runtime_days is not None and runtime_days < 0 else 0,
                    "Flag_BefundlaufzeitGroesser7Tage": 1 if runtime_days is not None and runtime_days > 7 else 0,
                    "Flag_BefundlaufzeitGroesser14Tage": 1 if runtime_days is not None and runtime_days > 14 else 0,
                    "Flag_FehlendesAuftragsdatum": 0 if order_date else 1,
                    "Flag_FehlendesBefunddatum": 0 if report_date else 1,
                }
                labor_next += 1

            material_standard, material_mapped = standardize_material(material_raw, material_map)
            material_counts[material_standard] += 1
            if material_raw and not material_mapped:
                unmapped_materials[material_raw] += 1
            anatomy = derive_anatomy(material_raw, localization_raw, info_raw, comment_raw)
            probe_source_key = (
                f"{labor_source_key}:probe:{stable_hash(material_raw, localization_raw, info_raw, salt=args.salt)}"
            )
            if probe_source_key not in probes:
                probes[probe_source_key] = {
                    "ProbeID": probe_next,
                    "SourceKey": probe_source_key,
                    "LaborauftragID": labor_orders[labor_source_key]["LaborauftragID"],
                    "LaborauftragSourceKey": labor_source_key,
                    "Material_Raw": material_raw,
                    "Material_Standard": material_standard,
                    "Probenbeschreibung_Raw": info_raw,
                    "Lokalisation_Raw": localization_raw,
                    **anatomy,
                    "PJI_Relevanz_Vorlaeufig": pji_relevance(
                        material_standard,
                        anatomy["Gelenk_Abgeleitet"],
                        anatomy["Implantatbezug_Abgeleitet"],
                        material_raw,
                        localization_raw,
                        info_raw,
                        comment_raw,
                    ),
                    "AnzahlErreger": 0,
                    "AnzahlKulturellPositiveErreger": 0,
                    "Polymikrobiell": 0,
                    "NurNiedrigVirulenteErreger": 0,
                    "MindestensEinHochvirulenterErreger": 0,
                    "MindestensEinResistenterErreger": 0,
                    "AntibiogrammVorhanden": 0,
                    "NurAnreicherungsnachweis": 0,
                    "GelenkFallID": None,
                    "EpisodeID": None,
                    "AufenthaltID": None,
                    "EingriffID": None,
                    "OPDatum": None,
                    "TageDifferenzZuOP": None,
                    "Zuordnungssicherheit": None,
                }
                probe_next += 1

            organism, organism_mapped = standardize_organism(pathogen_raw, organism_map)
            organism["Erreger_Raw"] = pathogen_raw
            if pathogen_raw and not organism_mapped:
                unmapped_organisms[pathogen_raw] += 1
            method = infer_method(material_standard, material_raw, localization_raw, info_raw, status_raw, comment_raw)
            flags = infer_result_flags(organism, method, pathogen_raw, status_raw, result_value_raw, comment_raw)
            quantity = quantity_fields(result_value_raw, comment_raw)
            isolate_source_key = (
                f"{probe_source_key}:isolate:{stable_hash(pathogen_raw, status_raw, result_value_raw, subtype_raw, comment_raw, salt=args.salt)}"
            )
            if isolate_source_key not in isolates:
                isolates[isolate_source_key] = {
                    "IsolatID": isolate_next,
                    "SourceKey": isolate_source_key,
                    "ProbeID": probes[probe_source_key]["ProbeID"],
                    "ProbeSourceKey": probe_source_key,
                    "LaborauftragID": labor_orders[labor_source_key]["LaborauftragID"],
                    "LaborauftragSourceKey": labor_source_key,
                    "Erreger_Raw": pathogen_raw,
                    "Erreger_Standard": organism["Erreger_Standard"],
                    "Erregergruppe": organism["Erregergruppe"],
                    "Gramstatus": organism["Gramstatus"],
                    "Virulenzgruppe": organism["Virulenzgruppe"],
                    "Resistenzmarker_Raw": resistance_marker_raw,
                    "ResistenzmarkerDatum": as_iso_date(first_value(row, col(config, "resistance_marker_date"))),
                    "Subtyp_Raw": subtype_raw,
                    "Ergebnisstatus_Raw": status_raw,
                    "Messwert_Raw": result_value_raw,
                    "Kommentar_Raw": comment_raw,
                    **flags,
                    "Nachweismethode": method,
                    **quantity,
                }
                isolate_next += 1
            isolate_by_probe[probe_source_key].add(isolate_source_key)

            ab_source_keys: list[str] = []
            for ab_column in ab_columns:
                antibiotic_columns_seen[ab_column] += 1
                result_raw = raw_text(row.get(ab_column))
                if not args.include_empty_antibiogram and (
                    result_raw is None or result_raw.lower().strip() in EMPTY_RESULTS
                ):
                    continue
                antibiotic_raw = antibiotic_from_header(ab_column)
                ab_key = normalized_key(antibiotic_raw)
                mapped_ab = antibiotic_map.get(ab_key)
                if not mapped_ab:
                    unmapped_antibiotics[antibiotic_raw] += 1
                sir, mic_raw = normalize_sir(result_raw)
                if sir == "UNKLAR":
                    invalid_sir[f"{ab_column}: {result_raw}"] += 1
                ab_source_key = f"{isolate_source_key}:ab:{ab_column}:{normalized_key(result_raw)}"
                if ab_source_key not in antibiograms:
                    antibiograms[ab_source_key] = {
                        "AntibiogrammID": antibiogram_next,
                        "SourceKey": ab_source_key,
                        "IsolatID": isolates[isolate_source_key]["IsolatID"],
                        "IsolatSourceKey": isolate_source_key,
                        "Antibiotikum_Raw": antibiotic_raw,
                        "Antibiotikum_Standard": (mapped_ab or {}).get("Antibiotikum_Standard") or antibiotic_raw,
                        "Antibiotikaklasse": (mapped_ab or {}).get("Antibiotikaklasse"),
                        "Spezialtest": (mapped_ab or {}).get("Spezialtest"),
                        "Ergebnis_Raw": result_raw,
                        "SIR_Standard": sir,
                        "MIC_Raw": mic_raw,
                        "Methode_Raw": None,
                        "Breakpointsystem": None,
                        "Testdatum": report_date,
                        "AntibiotikumSpalte": ab_column,
                    }
                    antibiogram_next += 1
                antibiogram_by_isolate[isolate_source_key].add(ab_source_key)
                ab_source_keys.append(ab_source_key)

            raw_links.append(
                {
                    "SheetName": sheet_name,
                    "RowNumber": excel_row_number,
                    "RawJSON": row_to_json(mask_identifier_row(row, include_identifiers=args.include_identifiers, salt=args.salt)),
                    "LaborauftragSourceKey": labor_source_key,
                    "ProbeSourceKey": probe_source_key,
                    "IsolatSourceKey": isolate_source_key,
                    "AntibiogrammSourceKeys": ";".join(ab_source_keys),
                }
            )
        if args.limit and counters["raw_rows"] >= args.limit:
            break

    for probe_source_key, isolate_keys in isolate_by_probe.items():
        probe = probes[probe_source_key]
        clinical_positive = []
        all_positive = []
        for isolate_key in isolate_keys:
            isolate = isolates[isolate_key]
            if isolate["Positiv"]:
                all_positive.append(isolate)
            if (
                isolate["Positiv"]
                and isolate["Erregergruppe"] not in {"NICHT_ERREGERBEFUND", "UNKLAR"}
                and isolate["Nachweismethode"] not in {"SCREENING", "PCR_PANEL"}
            ):
                clinical_positive.append(isolate)
        distinct = {item["Erreger_Standard"] for item in clinical_positive if item["Erreger_Standard"]}
        probe["AnzahlErreger"] = len({isolates[key]["Erreger_Standard"] for key in isolate_keys if isolates[key]["Erreger_Standard"]})
        probe["AnzahlKulturellPositiveErreger"] = len(distinct)
        probe["Polymikrobiell"] = 1 if len(distinct) >= 2 else 0
        probe["NurNiedrigVirulenteErreger"] = 1 if clinical_positive and all(
            item["Virulenzgruppe"] == "NIEDRIGVIRULENT" for item in clinical_positive
        ) else 0
        probe["MindestensEinHochvirulenterErreger"] = 1 if any(
            item["Virulenzgruppe"] == "HOCHVIRULENT" for item in clinical_positive
        ) else 0
        probe["AntibiogrammVorhanden"] = 1 if any(antibiogram_by_isolate.get(key) for key in isolate_keys) else 0
        probe["NurAnreicherungsnachweis"] = 1 if all_positive and all(item["NurAnreicherung"] for item in all_positive) else 0
        probe["MindestensEinResistenterErreger"] = 1 if any(
            antibiograms[ab_key]["SIR_Standard"] == "R"
            for isolate_key in isolate_keys
            for ab_key in antibiogram_by_isolate.get(isolate_key, set())
        ) else 0

    duplicate_labs = {key: len(value) for key, value in lab_patient_assignments.items() if len(value) > 1}
    duplicate_raw_rows = sum(count - 1 for count in duplicate_raw_counter.values() if count > 1)
    same_standard_ab = defaultdict(set)
    for row in antibiograms.values():
        same_standard_ab[row["Antibiotikum_Standard"]].add(row["Antibiotikum_Raw"])
    duplicate_ab_names = [
        {"Antibiotikum_Standard": standard, "RawSpalten": "; ".join(sorted(raw_values)), "AnzahlRaw": len(raw_values)}
        for standard, raw_values in sorted(same_standard_ab.items())
        if len(raw_values) > 1
    ]

    return {
        "patients": list(patients.values()),
        "labor_orders": list(labor_orders.values()),
        "probes": list(probes.values()),
        "isolates": list(isolates.values()),
        "antibiograms": list(antibiograms.values()),
        "raw_links": raw_links,
        "material_map_rows": material_map_rows,
        "organism_map_rows": organism_map_rows,
        "antibiotic_map_rows": antibiotic_map_rows,
        "unmapped_materials": [{"Material_Raw": key, "Anzahl": value} for key, value in unmapped_materials.most_common()],
        "unmapped_organisms": [{"Erreger_Raw": key, "Anzahl": value} for key, value in unmapped_organisms.most_common()],
        "unmapped_antibiotics": [{"Antibiotikum_Raw": key, "Anzahl": value} for key, value in unmapped_antibiotics.most_common()],
        "invalid_sir": [{"Wert": key, "Anzahl": value} for key, value in invalid_sir.most_common()],
        "duplicate_labs": [{"LaborauftragSourceKey": key, "Patienten": value} for key, value in duplicate_labs.items()],
        "duplicate_ab_names": duplicate_ab_names,
        "material_counts": [{"Material_Standard": key, "Anzahl": value} for key, value in material_counts.most_common()],
        "summary": [
            ("Quelldatei", str(args.file)),
            ("Tabellenblaetter", ", ".join(sheets)),
            ("Rohzeilen", counters["raw_rows"]),
            ("Eindeutige Patienten", len(patients)),
            ("Eindeutige Laborauftraege", len(labor_orders)),
            ("Proben", len(probes)),
            ("Isolate/Befundzeilen", len(isolates)),
            ("Antibiogramm-Eintraege", len(antibiograms)),
            ("Positive Isolate", sum(1 for row in isolates.values() if row["Positiv"])),
            ("Negative Isolate", sum(1 for row in isolates.values() if row["Negativ"])),
            ("Fehlende Laborauftragsnummern", counters["missing_labor_number"]),
            ("Identische Duplikatzeilen", duplicate_raw_rows),
            ("Widerspruechliche Laborauftrag-Patient-Zuordnungen", len(duplicate_labs)),
            ("Unbekannte Materialien", len(unmapped_materials)),
            ("Unbekannte Erreger", len(unmapped_organisms)),
            ("Unbekannte Antibiotika", len(unmapped_antibiotics)),
            ("Ungueltige S/I/R-Werte", sum(invalid_sir.values())),
        ],
    }


def insert_or_ignore(con: Any, table: str, row: dict[str, Any], columns: list[str]) -> None:
    placeholders = ", ".join("?" for _ in columns)
    column_sql = ", ".join(columns)
    con.execute(
        f"INSERT OR IGNORE INTO {table} ({column_sql}) VALUES ({placeholders})",
        [row.get(column) for column in columns],
    )


def fetch_id_map(con: Any, table: str, key_column: str, id_column: str) -> dict[str, int]:
    return {str(row[key_column]): int(row[id_column]) for row in con.execute(f"SELECT {key_column}, {id_column} FROM {table}")}


def write_to_db(args: argparse.Namespace, result: dict[str, Any]) -> None:
    with connect_db(args.db) as con:
        source_id = get_or_create_source(con, args.file, "HYBASE")
        batch_id = start_import_batch(con, source_id, "hybase_structured")
        try:
            for row in result["material_map_rows"]:
                insert_or_ignore(
                    con,
                    "mapMaterial",
                    row,
                    ["Material_Raw", "Material_Standard", "Kommentar", "Aktiv"],
                )
            for row in result["organism_map_rows"]:
                insert_or_ignore(
                    con,
                    "mapErreger",
                    row,
                    ["Erreger_Raw", "Erreger_Standard", "Erregergruppe", "Gramstatus", "Virulenzgruppe", "Kommentar", "Aktiv"],
                )
            for row in result["antibiotic_map_rows"]:
                insert_or_ignore(
                    con,
                    "mapAntibiotikum",
                    row,
                    ["Antibiotikum_Raw", "Antibiotikum_Standard", "Antibiotikaklasse", "Spezialtest", "Kommentar", "Aktiv"],
                )

            for row in result["patients"]:
                con.execute(
                    """
                    INSERT OR IGNORE INTO tblPatient (
                        PatientHash, PatientenzahlSAP, HYBASEPatientID_Raw,
                        HYBASEPatientHash, Geburtsdatum, Geburtsjahr, Geschlecht, Quelle
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, 'HYBASE')
                    """,
                    (
                        row["PatientHash"],
                        row["HYBASEPatientID_Raw"],
                        row["HYBASEPatientID_Raw"],
                        row["PatientHash"],
                        row["Geburtsdatum"],
                        row["Geburtsjahr"],
                        row["Geschlecht"],
                    ),
                )
            patient_ids = {str(row["PatientHash"]): int(row["PatientID"]) for row in con.execute("SELECT PatientID, PatientHash FROM tblPatient")}

            for row in result["labor_orders"]:
                db_row = dict(row)
                db_row["PatientID"] = patient_ids.get(row["PatientHash"])
                insert_or_ignore(
                    con,
                    "tblLaborauftrag",
                    db_row,
                    [field for field in LABORAUFTRAG_FIELDS if field not in {"LaborauftragID", "PatientHash"}],
                )
            labor_ids = fetch_id_map(con, "tblLaborauftrag", "SourceKey", "LaborauftragID")

            for row in result["probes"]:
                db_row = dict(row)
                db_row["LaborauftragID"] = labor_ids.get(row["LaborauftragSourceKey"])
                insert_or_ignore(
                    con,
                    "tblProbe",
                    db_row,
                    [field for field in PROBE_FIELDS if field not in {"ProbeID", "LaborauftragSourceKey"}],
                )
            probe_ids = fetch_id_map(con, "tblProbe", "SourceKey", "ProbeID")

            for row in result["isolates"]:
                db_row = dict(row)
                db_row["ProbeID"] = probe_ids.get(row["ProbeSourceKey"])
                db_row["LaborauftragID"] = labor_ids.get(row["LaborauftragSourceKey"])
                insert_or_ignore(
                    con,
                    "tblIsolat",
                    db_row,
                    [field for field in ISOLAT_FIELDS if field not in {"IsolatID", "ProbeSourceKey", "LaborauftragSourceKey"}],
                )
            isolate_ids = fetch_id_map(con, "tblIsolat", "SourceKey", "IsolatID")

            for row in result["antibiograms"]:
                db_row = dict(row)
                db_row["IsolatID"] = isolate_ids.get(row["IsolatSourceKey"])
                insert_or_ignore(
                    con,
                    "tblHybaseAntibiogramm",
                    db_row,
                    [field for field in ANTIBIOGRAMM_FIELDS if field not in {"AntibiogrammID", "IsolatSourceKey"}],
                )
            antibiogram_ids = fetch_id_map(con, "tblHybaseAntibiogramm", "SourceKey", "HybaseAntibiogrammID")

            for link in result["raw_links"]:
                raw_id = con.execute(
                    """
                    INSERT INTO tblRawHybaseImport (
                        ImportBatchID, SheetName, RowNumber, LaborauftragID, ProbeID,
                        IsolatID, HybaseAntibiogrammID, RawJSON
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        batch_id,
                        link["SheetName"],
                        link["RowNumber"],
                        labor_ids.get(link["LaborauftragSourceKey"]),
                        probe_ids.get(link["ProbeSourceKey"]),
                        isolate_ids.get(link["IsolatSourceKey"]),
                        None,
                        link["RawJSON"],
                    ),
                ).lastrowid
                for ab_key in filter(None, str(link.get("AntibiogrammSourceKeys") or "").split(";")):
                    ab_id = antibiogram_ids.get(ab_key)
                    if ab_id:
                        con.execute(
                            """
                            INSERT INTO tblRawHybaseIsolatAntibiogrammLink (
                                RawHybaseImportID, HybaseAntibiogrammID, AntibiotikumSpalte
                            )
                            VALUES (?, ?, ?)
                            """,
                            (raw_id, ab_id, ab_key.split(":ab:", 1)[-1].split(":", 1)[0]),
                        )

            finish_import_batch(
                con,
                batch_id,
                int(dict(result["summary"]).get("Rohzeilen", 0)),
                note=f"structured_labor_orders={len(result['labor_orders'])}; probes={len(result['probes'])}; isolates={len(result['isolates'])}; antibiograms={len(result['antibiograms'])}",
            )
            con.commit()
        except Exception as exc:
            finish_import_batch(con, batch_id, 0, status="failed", note=str(exc))
            con.commit()
            raise


def write_outputs(args: argparse.Namespace, result: dict[str, Any]) -> None:
    args.output_dir.mkdir(parents=True, exist_ok=True)
    write_csv(args.output_dir / "tblPatient.csv", result["patients"], PATIENT_FIELDS)
    write_csv(args.output_dir / "tblLaborauftrag.csv", result["labor_orders"], LABORAUFTRAG_FIELDS)
    write_csv(args.output_dir / "tblProbe.csv", result["probes"], PROBE_FIELDS)
    write_csv(args.output_dir / "tblIsolat.csv", result["isolates"], ISOLAT_FIELDS)
    write_csv(args.output_dir / "tblAntibiogramm.csv", result["antibiograms"], ANTIBIOGRAMM_FIELDS)
    write_csv(args.output_dir / "_raw_links.csv", result["raw_links"], RAW_LINK_FIELDS)
    write_csv(args.output_dir / "unmapped_materials.csv", result["unmapped_materials"], ["Material_Raw", "Anzahl"])
    write_csv(args.output_dir / "unmapped_organisms.csv", result["unmapped_organisms"], ["Erreger_Raw", "Anzahl"])
    write_csv(args.output_dir / "unmapped_antibiotics.csv", result["unmapped_antibiotics"], ["Antibiotikum_Raw", "Anzahl"])
    copy_mapping_csv(
        args.output_dir / "mapMaterial.csv",
        result["material_map_rows"],
        ["Material_Raw", "Material_Standard", "Kommentar", "Aktiv"],
    )
    copy_mapping_csv(
        args.output_dir / "mapErreger.csv",
        result["organism_map_rows"],
        ["Erreger_Raw", "Erreger_Standard", "Erregergruppe", "Gramstatus", "Virulenzgruppe", "Kommentar", "Aktiv"],
    )
    copy_mapping_csv(
        args.output_dir / "mapAntibiotikum.csv",
        result["antibiotic_map_rows"],
        ["Antibiotikum_Raw", "Antibiotikum_Standard", "Antibiotikaklasse", "Spezialtest", "Kommentar", "Aktiv"],
    )
    make_quality_report(
        args.output_dir / "quality_report.xlsx",
        result["summary"],
        {
            "Materialgruppen": result["material_counts"],
            "UnmappedMaterial": result["unmapped_materials"],
            "UnmappedErreger": result["unmapped_organisms"],
            "UnmappedAntibiotika": result["unmapped_antibiotics"],
            "InvalidSIR": result["invalid_sir"],
            "DoppelteAB": result["duplicate_ab_names"],
            "LaborPatientKonflikte": result["duplicate_labs"],
        },
    )
    if not args.no_database_import_xlsx:
        write_database_import(
            args.output_dir / "database_import.xlsx",
            {
                "tblPatient": (result["patients"], PATIENT_FIELDS),
                "tblLaborauftrag": (result["labor_orders"], LABORAUFTRAG_FIELDS),
                "tblProbe": (result["probes"], PROBE_FIELDS),
                "tblIsolat": (result["isolates"], ISOLAT_FIELDS),
                "tblAntibiogramm": (result["antibiograms"], ANTIBIOGRAMM_FIELDS),
            },
        )


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Strukturiert einen HYBASE-Mikrobiologieexport in Laborauftrag, Probe, Isolat und Antibiogramm."
    )
    parser.add_argument("file", type=Path)
    parser.add_argument("--config", type=Path, default=Path("config/hybase/hybase_columns.json"))
    parser.add_argument("--output-dir", type=Path, default=Path("outputs/hybase_structured"))
    parser.add_argument("--db", type=Path, default=Path("outputs/pji_register_orthodok_pseudonymized.sqlite"))
    parser.add_argument("--sheets", help="Kommagetrennte Excel-Blaetter. Standard kommt aus config.")
    parser.add_argument("--limit", type=int, help="Maximale Rohzeilen fuer sichere Tests.")
    parser.add_argument("--include-empty-antibiogram", action="store_true")
    parser.add_argument("--include-identifiers", action="store_true")
    parser.add_argument("--write-db", action="store_true", help="Schreibt die strukturierten Tabellen zusaetzlich in SQLite.")
    parser.add_argument("--no-database-import-xlsx", action="store_true")
    parser.add_argument("--salt", default="pji-register-local-v1")
    parser.add_argument("--log-level", default="INFO")
    args = parser.parse_args()

    logging.basicConfig(level=getattr(logging, args.log_level.upper(), logging.INFO), format="%(levelname)s: %(message)s")
    result = process_hybase(args)
    write_outputs(args, result)
    if args.write_db:
        write_to_db(args, result)
    for key, value in result["summary"]:
        print(f"{key}: {value}")
    print(f"Outputs: {args.output_dir}")
    if args.write_db:
        print(f"SQLite aktualisiert: {args.db}")


if __name__ == "__main__":
    main()
